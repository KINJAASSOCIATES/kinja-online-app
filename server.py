#!/usr/bin/env python3
"""
Kinja Associates - server
-------------------------
Ek hi file: app (index.html) dikhata hai, data save/sync karta hai,
password (login) lagata hai, roz ka backup rakhta hai aur Excel banata hai.

Local chalane ke liye:
    python server.py
    (phir browser me http://localhost:8000 kholein)

Online (hosting) par ye settings (Environment Variables) lagti hain:
    APP_USER      login ka naam          (default: admin)
    APP_PASSWORD  login ka password      (ZARURI - online karte waqt mazboot rakhein)
    DATA_DIR      data kahan save ho     (default: ./data)
    PORT          port number            (hosting khud deti hai, default 8000)

Ye file WSGI app bhi hai (`application`), isliye gunicorn / PythonAnywhere par bhi chalti hai.
"""
import os, sys, json, time, hmac, base64, io, shutil, threading, datetime
from urllib.parse import urlparse, parse_qs

try:
    import fcntl  # Linux / Mac: kayi workers ke beech safe locking
except ImportError:  # Windows
    fcntl = None

BASE = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.environ.get("DATA_DIR") or os.path.join(BASE, "data")
BACKUP_DIR = os.path.join(DATA_DIR, "backups")
DATA_FILE = os.path.join(DATA_DIR, "data.json")
LOCK_FILE = os.path.join(DATA_DIR, ".lock")
INDEX_FILE = os.path.join(BASE, "index.html")
APP_USER = os.environ.get("APP_USER", "admin")
APP_PASSWORD = os.environ.get("APP_PASSWORD", "")
KEEP_BACKUPS = 60            # itne din ke backup rakhe jayenge
MAX_BODY = 25 * 1024 * 1024  # 25 MB

os.makedirs(BACKUP_DIR, exist_ok=True)
_tlock = threading.Lock()


class _Lock:
    """Thread + process lock, taaki do log ek saath save karein to data kharab na ho."""
    def __enter__(self):
        _tlock.acquire()
        self.f = open(LOCK_FILE, "a+")
        if fcntl:
            fcntl.flock(self.f, fcntl.LOCK_EX)
        return self

    def __exit__(self, *a):
        if fcntl:
            fcntl.flock(self.f, fcntl.LOCK_UN)
        self.f.close()
        _tlock.release()


def read_store():
    try:
        with open(DATA_FILE, "r", encoding="utf-8") as f:
            s = json.load(f)
        return {"ver": int(s.get("ver", 0)), "data": s.get("data")}
    except (FileNotFoundError, ValueError):
        return {"ver": 0, "data": None}


def write_store(store):
    tmp = DATA_FILE + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(store, f, ensure_ascii=False)
        f.flush()
        os.fsync(f.fileno())
    os.replace(tmp, DATA_FILE)  # atomic: beech me band hone par bhi file kharab nahi hoti
    # roz ka backup (us din ki latest copy)
    try:
        day = datetime.date.today().isoformat()
        shutil.copyfile(DATA_FILE, os.path.join(BACKUP_DIR, "data-%s.json" % day))
        files = sorted(x for x in os.listdir(BACKUP_DIR) if x.startswith("data-"))
        for old in files[:-KEEP_BACKUPS]:
            os.remove(os.path.join(BACKUP_DIR, old))
    except OSError:
        pass


def make_xlsx(payload):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    headers = payload.get("headers") or []
    rows = payload.get("rows") or []
    ws.append([str(payload.get("title", ""))])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(headers)
    hr = ws.max_row
    num_cols = set()
    for r in rows:
        ws.append(r)
        for j, v in enumerate(r):
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                num_cols.add(j)
    for j, _ in enumerate(headers):
        c = ws.cell(row=hr, column=j + 1)
        c.font = Font(bold=True)
        if j in num_cols:
            c.alignment = Alignment(horizontal="right")  # amount heading ke neeche
    for r in ws.iter_rows(min_row=hr + 1):
        for c in r:
            if isinstance(c.value, (int, float)):
                c.number_format = "#,##0.00"
                c.alignment = Alignment(horizontal="right")
    if rows:
        for c in ws[hr + len(rows)]:  # last (total) row bold
            c.font = Font(bold=True)
    for col in ws.columns:
        w = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(max(w + 2, 10), 45)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _authorized(environ):
    if not APP_PASSWORD:
        return True
    h = environ.get("HTTP_AUTHORIZATION", "")
    if not h.startswith("Basic "):
        return False
    try:
        u, _, p = base64.b64decode(h[6:]).decode("utf-8").partition(":")
    except Exception:
        return False
    ok_u = hmac.compare_digest(u.encode(), APP_USER.encode())
    ok_p = hmac.compare_digest(p.encode(), APP_PASSWORD.encode())
    return ok_u and ok_p


def _resp(start, code, body=b"", ctype="application/json", extra=None):
    if isinstance(body, (dict, list)):
        body = json.dumps(body, ensure_ascii=False).encode("utf-8")
    elif isinstance(body, str):
        body = body.encode("utf-8")
    hd = [("Content-Type", ctype), ("Content-Length", str(len(body))),
          ("Cache-Control", "no-store"), ("X-Content-Type-Options", "nosniff"),
          ("X-Frame-Options", "DENY")]
    hd += extra or []
    start(code, hd)
    return [body]


def application(environ, start_response):
    path = urlparse(environ.get("PATH_INFO", "/") or "/").path
    method = environ.get("REQUEST_METHOD", "GET").upper()

    if not _authorized(environ):
        time.sleep(1)  # galat password try karne walon ko dheema karta hai
        return _resp(start_response, "401 Unauthorized", "Login zaroori hai", "text/plain; charset=utf-8",
                     [("WWW-Authenticate", 'Basic realm="Kinja Associates", charset="UTF-8"')])

    try:
        if method == "GET" and path in ("/", "/index.html"):
            with open(INDEX_FILE, "rb") as f:
                return _resp(start_response, "200 OK", f.read(), "text/html; charset=utf-8")

        if method == "GET" and path == "/api/ver":
            return _resp(start_response, "200 OK", {"ver": read_store()["ver"]})

        if method == "GET" and path == "/api/data":
            with _Lock():
                return _resp(start_response, "200 OK", read_store())

        if method == "POST" and path == "/api/data":
            n = int(environ.get("CONTENT_LENGTH") or 0)
            if n > MAX_BODY:
                return _resp(start_response, "413 Request Entity Too Large", {"error": "too_big"})
            body = json.loads(environ["wsgi.input"].read(n) or b"{}")
            with _Lock():
                cur = read_store()
                if int(body.get("ver", -1)) != cur["ver"]:
                    return _resp(start_response, "409 Conflict", {"error": "conflict", "ver": cur["ver"]})
                store = {"ver": cur["ver"] + 1, "data": body.get("data")}
                write_store(store)
            return _resp(start_response, "200 OK", {"ver": store["ver"]})

        if method == "POST" and path == "/api/xlsx":
            n = int(environ.get("CONTENT_LENGTH") or 0)
            if n > MAX_BODY:
                return _resp(start_response, "413 Request Entity Too Large", {"error": "too_big"})
            payload = json.loads(environ["wsgi.input"].read(n) or b"{}")
            data = make_xlsx(payload)
            return _resp(start_response, "200 OK", data,
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        if method == "GET" and path == "/api/gst":
            # GST API ke bina naam/address apne aap nahi aata (app khud manual option dikhata hai)
            return _resp(start_response, "200 OK", {"error": "not_configured"})

        if method == "GET" and path == "/api/backup":
            # poora data download: browser me  /api/backup  kholein
            with _Lock():
                raw = json.dumps(read_store(), ensure_ascii=False)
            name = "kinja-backup-%s.json" % datetime.date.today().isoformat()
            return _resp(start_response, "200 OK", raw, "application/json",
                         [("Content-Disposition", 'attachment; filename="%s"' % name)])

        return _resp(start_response, "404 Not Found", {"error": "not_found"})
    except ValueError:
        return _resp(start_response, "400 Bad Request", {"error": "bad_request"})
    except Exception as e:  # server band na ho, error dikha do
        sys.stderr.write("ERROR %r\n" % (e,))
        return _resp(start_response, "500 Internal Server Error", {"error": "server_error"})


if __name__ == "__main__":
    from wsgiref.simple_server import make_server, WSGIServer, WSGIRequestHandler
    from socketserver import ThreadingMixIn

    class Threaded(ThreadingMixIn, WSGIServer):
        daemon_threads = True

    class Quiet(WSGIRequestHandler):
        def log_message(self, *a):
            pass

    port = int(os.environ.get("PORT", "8000"))
    host = os.environ.get("HOST", "0.0.0.0")
    if not APP_PASSWORD:
        print("WARNING: APP_PASSWORD set nahi hai - bina login ke chal raha hai. "
              "Online karne se pehle password zaroor lagayein.")
    print("Server chal raha hai: http://localhost:%d   (band karne ke liye Ctrl+C)" % port)
    make_server(host, port, application, Threaded, Quiet).serve_forever()
